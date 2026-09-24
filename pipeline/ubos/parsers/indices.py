"""Price indices beyond the CPI: house prices (RPPI), construction inputs (CIPI),
producer prices (PPI). Each latest release carries the full history, so only the
newest workbook is parsed. We read index levels, derive annual % changes
ourselves, and check them against what UBOS printed (or, where UBOS prints no
rates, against the index's own weighting).

RPPI  sheet "tables": FY + quarter rows; Wakiso, Kampala&Makindye, Nakawa,
      Kawempe&Rubaga, Headline; UBOS quarterly and annual % change columns.
      Covers greater Kampala only.
CIPI  sheets "ALL CONSTRUCTION", "41 - BUILDINGS", "42 - CIVIL ENGINEERING",
      "43 - SPECIALISED CONSTRUCTION": product rows x monthly columns, with
      "Materials" and "OVERALL CONSTRUCTION INDEX" aggregate rows.
PPI   sheet "INDICES PER DIVISION" (index levels) and
      "ANNUAL INFLATION PER DIVISION" (UBOS's annual % change, used as a check).
"""

from __future__ import annotations

import re
from datetime import date, datetime

import openpyxl

from .wide import num

QUARTER = re.compile(r"^Q[1-4]$")


def _rows(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]


def _month(v) -> str | None:
    if isinstance(v, datetime | date):
        return f"{v.year:04d}-{v.month:02d}"
    if isinstance(v, str) and re.match(r"^\d{4}-\d{2}-\d{2}", v):
        return v[:7]
    return None


def _yoy(values: list, lag: int) -> list:
    return [None] * lag + [
        round((b / a - 1) * 100, 2) if a and b else None for a, b in zip(values, values[lag:])
    ]


def _check_rates(ours: list, theirs: list, labels: list, what: str, tol: float = 0.15):
    checked = 0
    for lab, a, b in zip(labels, ours, theirs):
        if a is None or b is None:
            continue
        checked += 1
        if abs(a - b) > tol:
            raise ValueError(f"{what}: derived {a}% != UBOS {b}% for {lab}")
    if checked == 0:
        raise ValueError(f"{what}: nothing to check against UBOS's published rates")


# ---- house prices ------------------------------------------------------------

RPPI_AREAS = {
    "Wakiso": "Wakiso",
    "Kampala&Makindye": "Kampala Central & Makindye",
    "Nakawa": "Nakawa",
    "Kawempe&Rubage": "Kawempe & Rubaga",
    "Headline": "Greater Kampala (headline)",
}


def parse_rppi(path) -> dict:
    rows = _rows(path, "tables")
    hdr = next(i for i, r in enumerate(rows) if any(str(c).strip() == "Headline" for c in r[:8] if c))
    cols = {str(c).strip(): j for j, c in enumerate(rows[hdr][:9]) if c}
    annual_col = next(j for name, j in cols.items() if name.lower().startswith("annual"))
    labels, fy = [], None
    series = {v: [] for v in RPPI_AREAS.values()}
    published = []
    for r in rows[hdr + 1 :]:
        if r[0] and re.match(r"^\d{4}/\d{2}$", str(r[0]).strip()):
            fy = str(r[0]).strip()
        q = str(r[1]).strip() if r[1] else ""
        if not (fy and QUARTER.match(q)):
            continue
        labels.append(f"{fy} {q}")
        for src, name in RPPI_AREAS.items():
            series[name].append(num(r[cols[src]]))
        published.append(num(r[annual_col]))
    headline = series[RPPI_AREAS["Headline"]]
    yoy = {name: _yoy(v, 4) for name, v in series.items()}
    _check_rates(yoy[RPPI_AREAS["Headline"]], published, labels, "RPPI headline annual change", tol=0.2)
    return {"quarters": labels, "index": series, "yoy": yoy, "base": "2015/16 = 100", "headline": headline}


# ---- construction input prices -----------------------------------------------

CIPI_SHEETS = {
    "ALL CONSTRUCTION": "All construction",
    "41 - BUILDINGS": "Buildings",
    "42 - CIVIL ENGINEERING": "Roads & civil works",
    "43 - SPECIALISED CONSTRUCTION": "Specialised works",
}


def _monthly_block(rows, label_col: int, weight_col: int):
    """Find the header row with month columns; return (months, [(label, weight, values)])."""
    for h, r in enumerate(rows):
        months = [(j, _month(c)) for j, c in enumerate(r) if _month(c)]
        if len(months) >= 12:
            break
    else:
        raise ValueError("no month header")
    out = []
    for r in rows[h + 1 :]:
        label = " ".join(str(r[label_col]).split()) if r[label_col] else ""
        if not label:
            continue
        vals = [num(r[j]) for j, _ in months]
        if any(v is not None for v in vals):
            out.append((label, num(r[weight_col]), vals))
    return [m for _, m in months], out


def parse_cipi(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = {ws.title.strip().upper(): ws.title for ws in wb.worksheets}
    kinds, months = {}, None
    for sheet, nice in CIPI_SHEETS.items():
        rows = [list(r) for r in wb[names[sheet]].iter_rows(values_only=True)]
        m, items = _monthly_block(rows, label_col=1, weight_col=2)
        if months is None:
            months = m
        elif m != months:
            raise ValueError(f"CIPI: {sheet} months differ from the other sheets")
        overall = next((v for lab, _, v in items if lab.upper().startswith("OVERALL")), None)
        if overall is None:
            raise ValueError(f"CIPI: no overall index row in {sheet}")
        kinds[nice] = {"overall": overall, "items": {lab: {"weight": w, "index": v} for lab, w, v in items if not lab.upper().startswith("OVERALL")}}

    # Check: in "All construction", Materials + Utilities + Services + Labour
    # reproduce the overall index (weighted), and their weights sum to 1000.
    allc = kinds["All construction"]["items"]
    parts = ["Materials", "Utilities", "Services", "Labour"]
    w = [allc[p]["weight"] for p in parts]
    if abs(sum(w) - 1000) > 1:
        raise ValueError(f"CIPI: component weights sum to {sum(w)}, not 1000")
    overall = kinds["All construction"]["overall"]
    for i, month in enumerate(months):
        est = sum(allc[p]["weight"] * (allc[p]["index"][i] or 0) for p in parts) / 1000
        if overall[i] and abs(est - overall[i]) / overall[i] > 0.01:
            raise ValueError(f"CIPI: weighted components {est:.1f} != overall {overall[i]} in {month}")

    for k in kinds.values():
        k["overall_yoy"] = _yoy(k["overall"], 12)
        for it in k["items"].values():
            it["yoy"] = _yoy(it["index"], 12)
    return {"months": months, "kinds": kinds, "base": "2016/17 = 100"}


# ---- producer prices ------------------------------------------------------------

def parse_ppi(path) -> dict:
    idx_rows = _rows(path, "INDICES PER DIVISION")
    months, items = _monthly_block(idx_rows, label_col=1, weight_col=2)
    codes = {}
    for r in idx_rows:
        label = " ".join(str(r[1]).split()) if r[1] else ""
        if label and r[0] is not None:
            codes.setdefault(label, str(r[0]).strip())
    series, seen = {}, set()
    for label, weight, vals in items:
        if label in seen:  # "MANUFACTURING" appears in both tables on the sheet
            continue
        seen.add(label)
        series[label] = {"code": codes.get(label), "weight": weight, "index": vals, "yoy": _yoy(vals, 12)}

    head = next(k for k in series if k.upper().startswith("MANUFACTURING & UTILITIES"))
    # Check our annual rates against UBOS's own table for the headline index.
    rates = _rows(path, "ANNUAL INFLATION PER DIVISION")
    r_months, r_items = _monthly_block(rates, label_col=0, weight_col=2)
    theirs = dict(zip(r_months, next(v for lab, _, v in r_items if lab.lower().startswith("manufacturing & utilities"))))
    ours = dict(zip(months, series[head]["yoy"]))
    common = [m for m in r_months if m in ours]
    _check_rates([ours[m] for m in common], [theirs[m] for m in common], common, "PPI annual inflation", tol=0.2)
    return {"months": months, "headline": head, "series": series, "base": "2016/17 = 100"}
