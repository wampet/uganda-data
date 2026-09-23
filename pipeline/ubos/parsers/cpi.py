"""Consumer Price Index (CPI) workbook parser.

The monthly "CPI Excel Tables" workbook carries the whole history since July
2017 (base FY2016/17 = 100), so only the latest release needs parsing:

  sheet "Division"   headline (Grand Total), 13 COICOP divisions, 10 price
                     centres (towns), each followed by UBOS's own % changes
  sheet "Decomposed" Core / Food crops / Energy-fuel-utilities / Liquid fuels /
                     Non-core indices and ~380 individual items

We read index LEVELS only, locating blocks by their labels, then derive monthly
and annual % changes ourselves and check them against the rates UBOS printed.
A mismatch means the layout changed under us, so we fail loudly.
"""

from __future__ import annotations

import re
from datetime import date, datetime

import openpyxl

ITEM_CODE = re.compile(r"^(\d{2}(?:\.\d+){3,5})\s*(.+)$")


def _month(v) -> str | None:
    if isinstance(v, datetime | date):
        return f"{v.year:04d}-{v.month:02d}"
    return None


def _clean(label) -> str:
    return " ".join(str(label).split()) if label is not None else ""


def _num(v):
    return round(float(v), 4) if isinstance(v, int | float) else None


def _header_months(row, first_col: int) -> list[str]:
    months = [_month(v) for v in row[first_col:]]
    while months and months[-1] is None:
        months.pop()
    if not months or any(m is None for m in months):
        raise ValueError("CPI: unexpected header row (non-date column in month header)")
    return months


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def parse(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = {ws.title.strip().lower(): ws for ws in wb.worksheets}
    division = list(sheets["division"].iter_rows(values_only=True))
    decomposed = list(sheets["decomposed"].iter_rows(values_only=True))

    # ---- Division sheet: col0 code, col1 label, col2 weight, col3.. months
    months = _header_months(division[0], 3)
    n = len(months)
    series: dict[str, dict] = {}
    published: dict[str, float] = {}  # UBOS's own annual % (latest month), for checks

    def add(sid, name, group, weight, values, code=None):
        vals = [_num(v) for v in values[:n]]
        if sum(v is not None for v in vals) < n * 0.5:
            return
        series[sid] = {"name": name, "group": group, "weight": _num(weight), "code": code, "index": vals}

    # Blocks, top to bottom: division indices -> "Grand Total" -> UBOS % change
    # tables -> "Centre" -> centre indices -> more % change tables.
    state = "divisions"
    for row in division[1:]:
        code, label, weight, data = row[0], _clean(row[1]), row[2], row[3 : 3 + n]
        key = label.lower()

        if key == "grand total":
            add("headline", "All items (headline)", "headline", weight, data)
            state = "rates"
        elif key == "centre":
            state = "centres"
        elif state == "divisions" and isinstance(code, int | float) and label:
            add(f"div-{int(code):02d}", label, "division", weight, data, code=f"{int(code):02d}")
        elif state == "centres":
            if key.startswith(("annual", "monthly")):
                state = "done"
            elif label and isinstance(weight, int | float):
                add(f"centre-{_slug(label)}", label, "centre", weight, data)
        # First "Headline" row sits in UBOS's annual % table; keep it to check against.
        if key == "headline" and "headline_annual" not in published:
            published["headline_annual"] = _num(data[-1])

    # ---- Decomposed sheet: col0 label, col1 weight, col2.. months
    dmonths = _header_months(decomposed[0], 2)
    if dmonths[: len(months)] != months[: len(dmonths)]:
        raise ValueError("CPI: Division and Decomposed month headers disagree")

    aggregates = {
        "core index": ("core", "Core inflation (excludes food crops & fuel)"),
        "food crops and related items index": ("food-crops", "Food crops"),
        "energy fuel and utilities (efu) index": ("energy", "Energy, fuel & utilities"),
        "liquid energy fuels index": ("fuel", "Liquid fuels (petrol, diesel, paraffin)"),
        "non-core": ("non-core", "Non-core"),
    }
    item_group = "core"
    in_summary = False  # after "Summary" only aggregates (e.g. Non-core) remain
    for row in decomposed[1:]:
        label, weight, data = _clean(row[0]), row[1], row[2 : 2 + n]
        key = label.lower()
        if key in aggregates:
            sid, name = aggregates[key]
            if sid not in series and isinstance(weight, int | float):
                add(sid, name, "aggregate", weight, data)
            item_group = sid
            continue
        if key == "summary":
            in_summary = True
        if in_summary:
            continue
        m = ITEM_CODE.match(label)
        if m and isinstance(weight, int | float):
            code, name = m.group(1), m.group(2).strip()
            sid = f"item-{code.replace('.', '-')}"
            if sid not in series:
                add(sid, name, "item", weight, data, code=code)
                if sid in series:
                    series[sid]["basket"] = item_group

    for row in decomposed:
        if _clean(row[0]).lower() == "core" and row[1] is not None:
            published["core_annual"] = _num(row[2 + n - 1])

    _derive_rates(series)
    _check(series, published, months)

    return {
        "id": "cpi",
        "title": "Consumer Price Index",
        "base": "FY2016/17 = 100",
        "months": months,
        "series": series,
    }


def _derive_rates(series: dict) -> None:
    for s in series.values():
        idx = s["index"]
        s["mom"] = [None] + [
            round((b / a - 1) * 100, 2) if a and b else None for a, b in zip(idx, idx[1:])
        ]
        s["yoy"] = [None] * 12 + [
            round((b / a - 1) * 100, 2) if a and b else None for a, b in zip(idx, idx[12:])
        ]


def _check(series: dict, published: dict, months: list[str]) -> None:
    required = ["headline", "core"]
    for sid in required:
        if sid not in series:
            raise ValueError(f"CPI: required series '{sid}' not found — layout changed?")
    divisions = [s for s in series.values() if s["group"] == "division"]
    centres = [s for s in series.values() if s["group"] == "centre"]
    if len(divisions) != 13:
        raise ValueError(f"CPI: expected 13 divisions, found {len(divisions)}")
    if len(centres) < 8:
        raise ValueError(f"CPI: expected ~10 price centres, found {len(centres)}")

    for key, sid in (("headline_annual", "headline"), ("core_annual", "core")):
        want = published.get(key)
        got = series[sid]["yoy"][-1]
        if want is None:
            raise ValueError(f"CPI: could not find UBOS-published {key} to check against")
        if abs(want - got) > 0.05:
            raise ValueError(f"CPI: derived {sid} annual rate {got} != UBOS published {want} ({months[-1]})")
